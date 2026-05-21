"""Convert project architecture markdown to Word (.docx)."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_code_font(run) -> None:
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Consolas")
    rFonts.set(qn("w:hAnsi"), "Consolas")
    rPr.append(rFonts)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        set_code_font(run)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)


def add_table_from_md(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(col_count):
            cell = row[j] if j < len(row) else ""
            table.rows[i].cells[j].text = cell.strip()


def parse_inline(text: str) -> tuple[str, list[tuple[str, bool]]]:
    """Return plain segments with bold markers as (text, is_bold)."""
    parts: list[tuple[str, bool]] = []
    pattern = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")
    last = 0
    for m in pattern.finditer(text):
        if m.start() > last:
            parts.append((text[last : m.start()], False))
        if m.group(1):
            parts.append((m.group(1), True))
        else:
            parts.append((m.group(2), False))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], False))
    return text, parts


def add_paragraph_with_inline(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    _, parts = parse_inline(text)
    for segment, is_bold in parts:
        if not segment:
            continue
        run = p.add_run(segment)
        run.bold = is_bold
        if "`" in text and not is_bold and segment.isidentifier():
            set_code_font(run)


def convert_md_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table_from_md(doc, table_rows)
            table_rows = []
            doc.add_paragraph()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line.rstrip("\n"))
            i += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            if stripped.replace("|", "").replace("-", "").replace(":", "").strip() == "":
                i += 1
                continue
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_rows.append(cells)
            i += 1
            continue
        flush_table()

        if stripped in ("---", "***", "___"):
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            add_paragraph_with_inline(doc, stripped[2:], style="List Bullet")
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            add_paragraph_with_inline(doc, m.group(2), style="List Number")
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        add_paragraph_with_inline(doc, stripped)
        i += 1

    flush_table()
    if in_code and code_lines:
        add_code_block(doc, code_lines)

    title = docx_path.stem.replace("_", " ")
    doc.core_properties.title = title
    doc.core_properties.subject = "Logo Design Portal / Hawk Merchandising — Architecture Onboarding"
    doc.save(docx_path)
    print(f"Created: {docx_path}")


if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else base / "ARCHITECTURE_ONBOARDING.md"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else base / "ARCHITECTURE_ONBOARDING.docx"
    if not src.exists():
        print(f"Missing source: {src}", file=sys.stderr)
        sys.exit(1)
    convert_md_to_docx(src, dst)
