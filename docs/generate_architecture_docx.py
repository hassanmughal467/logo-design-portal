# -*- coding: utf-8 -*-
"""Generate ARCHITECTURE_ONBOARDING.docx for the Web Portal project."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUT = Path(__file__).parent / "ARCHITECTURE_ONBOARDING.docx"
OUT_FALLBACK = Path(__file__).parent / "ARCHITECTURE_ONBOARDING.generated.docx"


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def p(doc, text):
    doc.add_paragraph(text)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            t.rows[i + 1].cells[j].text = cell
    doc.add_paragraph()


def main():
    doc = Document()
    doc.core_properties.title = "Architecture Onboarding — Logo Design Portal"
    doc.core_properties.subject = "Hawk Merchandising Web Portal"

    h(doc, "Logo Design Portal — Architecture Onboarding", 0)
    p(doc, "Hawk Merchandising Web Portal | Last updated: May 2026")
    p(doc, "Source of truth: Backend/src/, Frontend/src/, docs/UPDATED_ARCHITECTURE_ONBOARDING.md")

    # 1
    h(doc, "1. Project Overview", 1)
    p(doc, "A B2B logo design operations portal: clients request work, admins assign designers, designers deliver previews, clients approve or revise, and the business bills clients and pays designers.")
    h(doc, "Main business purpose", 2)
    bullets(doc, [
        "Mediated design marketplace (clients never see designer identity).",
        "Order lifecycle from quote/creation through completion.",
        "Dual pricing: client charge (USD) vs designer payout (PKR-centric fields).",
        "Invoicing, payments (PayPal, Wise, bank), designer payout invoices.",
        "Admin analytics (revenue, churn, workflow) and client financial insights.",
    ])
    h(doc, "Main user types", 2)
    table(doc, ["Role", "Purpose"], [
        ["SuperAdmin", "Full access; permission grants; root admin protection"],
        ["Admin", "Operations; masked client PII"],
        ["Designer", "Assigned orders only; no client identity"],
        ["Client", "Own orders; sees Company Design Team instead of designer"],
    ])
    p(doc, "Roles: Backend/src/LogoDesignPortal.Domain/Enums/SystemRoles.cs")
    h(doc, "Core workflows", 2)
    bullets(doc, [
        "Register/login → HttpOnly JWT cookies + CSRF (or Bearer for tests).",
        "Optional quote → convert to order (QuotesController, /quotes UI).",
        "Create order → WaitingForAdminApproval.",
        "Admin assigns designer → InProgress (PriceApprovalPending when pricing changes).",
        "Preview → revision or approval → Completed → billing → invoice → payment.",
        "Designer payout invoice from completed orders.",
    ])
    p(doc, "State machine: Backend/src/LogoDesignPortal.Domain/OrderStatusStateMachine.cs")

    # 2
    h(doc, "2. Tech Stack", 1)
    table(doc, ["Layer", "Technologies"], [
        ["Frontend", "Angular 15.2, PrimeNG 15, RxJS, SignalR, Chart.js/ApexCharts, Playwright, Karma"],
        ["Backend", ".NET 8, ASP.NET Core, EF Core 8, JWT, AutoMapper, BCrypt, Hangfire, Serilog"],
        ["Database", "MySQL 8 (prod); SQLite (tests)"],
        ["Cache/Jobs", "Redis (optional in-memory fallback), Hangfire"],
        ["Hosting", "IIS in-process; SPA static dist"],
    ])
    p(doc, "Production API: https://api.hawkmerchandising.com | Admin UI: https://admin.hawkmerchandising.com")

    # 3
    h(doc, "3. Folder & Architecture Breakdown", 1)
    p(doc, "Modular monolith: single API + Angular SPA. Layer-based backend, feature-module frontend.")
    bullets(doc, [
        "Backend/src/LogoDesignPortal.Domain — entities, enums, state machine",
        "Backend/src/LogoDesignPortal.Application — services, DTOs, interfaces",
        "Backend/src/LogoDesignPortal.Infrastructure — EF, migrations, JWT, email",
        "Backend/src/LogoDesignPortal.API — controllers, middleware, hubs",
        "Frontend/src/app/core — auth, interceptors, guards",
        "Frontend/src/app/{feature} — lazy modules (orders, invoices, etc.)",
    ])
    p(doc, "WARNING: Ignore legacy Backend/LogoDesignPortal.API/ (not in solution).")

    # 4
    h(doc, "4. Entry Points & Application Flow", 1)
    p(doc, "Backend: Program.cs → DI → DatabaseInitializationHostedService → pipeline.")
    p(doc, "Frontend: main.ts → AppModule → AuthService restores localStorage → routing.")
    p(doc, "API pipeline: Swagger (Dev/Staging only) → ForwardedHeaders → SecurityHeaders → CORS → CorrelationId → SlowRequest → Serilog → Exception → HTTPS → AuthN → CsrfValidation → AuthZ → Hangfire → RateLimit → Controllers/Hub/Health.")
    p(doc, "Startup: DatabaseInitializationHostedService, FileStorageInitializer, ScalabilityServiceRegistration recurring jobs.")

    # 5
    h(doc, "5. Business Logic Analysis", 1)
    h(doc, "Order status transitions", 2)
    p(doc, "Enforced via OrderStatusStateMachine; workflow code should use OrderStatusTransitionHelper.")
    p(doc, "WaitingForAdminApproval → InProgress | PriceApprovalPending | Cancelled*")
    p(doc, "PriceApprovalPending → InProgress | WaitingForAdminApproval | Cancelled*")
    p(doc, "InProgress → PreviewDelivered | PriceApprovalPending | Cancelled*")
    p(doc, "PreviewDelivered → RevisionRequested | ClientApproved | Cancelled*")
    p(doc, "RevisionRequested → PreviewDelivered | InProgress | PriceApprovalPending | Cancelled*")
    p(doc, "ClientApproved → Completed → Refunded")
    h(doc, "Quotes (pre-order)", 2)
    p(doc, "api/quotes — client create, admin review/convert. Frontend: lazy quotes module. Migration: AddQuotesSystem.")
    h(doc, "Privacy / mediation rules", 2)
    table(doc, ["Viewer", "Rule", "File"], [
        ["Admin", "Client email/phone omitted", "OrderService.cs"],
        ["Designer", "Client info null on orders", "OrderService.cs"],
        ["Client", "Designer shown as Company Design Team", "OrderService.cs, FileService.cs"],
        ["Designer", "Client comments filtered", "CommentService.cs"],
    ])
    h(doc, "Production kill-switches", 2)
    p(doc, "ProductionSafety: DisableBillingGeneration, DisableDesignerPayout, DisableFileUploads, DisableInvoiceEditing. See docs/PRODUCTION-SAFETY-KILL-SWITCH.md")

    # 6
    h(doc, "6. Database & Data Models", 1)
    p(doc, "All entities extend BaseEntity (Id, audit, soft delete IsDeleted).")
    p(doc, "DbContext: Backend/src/LogoDesignPortal.Infrastructure/Persistence/ApplicationDbContext.cs")
    p(doc, "Migrations: Backend/src/LogoDesignPortal.Infrastructure/Migrations/ (~42 migrations)")
    bullets(doc, [
        "User → Role; User 1:1 ClientProfile or DesignerProfile",
        "LogoOrder → ClientProfile, optional DesignerProfile",
        "LogoOrder → LogoFile, OrderComment, OrderRevision, InvoiceOrder",
        "Invoice → Payment, InvoiceLog; DesignerInvoice → items/adjustments",
        "ClientLogoPricing, DesignerLogoPricing, Quote, Settings, AuditLog",
    ])
    p(doc, "Optimistic concurrency: LogoOrder.RowVersion, Invoice.RowVersion")

    # 7
    h(doc, "7. API Documentation", 1)
    p(doc, "Base: {apiUrl}/api | Auth: Authorization: Bearer {token} | JSON camelCase")
    table(doc, ["Area", "Route prefix", "Notes"], [
        ["Auth", "api/auth", "login, register, refresh, logout (clears cookies), password reset"],
        ["Orders", "api/orders", "CRUD, assign, status, batch preview, refund"],
        ["Quotes", "api/quotes", "Client quotes, admin convert to order"],
        ["Files", "api/files", "upload, download (visibility + IDOR checks)"],
        ["Billing", "api/billing", "Billing queue, client billing"],
        ["Invoices", "api/invoices", "generate, mark-paid, PDF"],
        ["Payments", "api/payments", "PayPal webhook (anonymous where applicable)"],
        ["Admin", "api/admin/*", "Financial, analytics, client intelligence"],
        ["Realtime", "hubs/notifications", "SignalR; access_token query or cookie"],
        ["Health", "health, health/ready", "DB schema, Hangfire, Redis"],
    ])
    p(doc, "24 controllers under Backend/src/LogoDesignPortal.API/Controllers/")

    # 8
    h(doc, "8. Authentication & Security", 1)
    h(doc, "HttpOnly cookie auth (recommended)", 2)
    bullets(doc, [
        "AuthCookieService sets ldp_access / ldp_refresh (HttpOnly) + ldp_csrf on login.",
        "CsrfValidationMiddleware: double-submit on mutating requests when cookie session is used.",
        "JwtBearer reads cookie in OnMessageReceived; Bearer header skips CSRF (integration tests).",
        "Frontend: useCookieAuth, CookieCredentialsInterceptor, CsrfInterceptor (X-XSRF-TOKEN).",
        "CORS AllowCredentials required for cross-origin SPA.",
    ])
    h(doc, "Authorization & hardening", 2)
    bullets(doc, [
        "JWT + refresh on User; BCrypt passwords; lockout after failed logins.",
        "Roles: [Authorize(Roles=...)]; [RequirePermission(...)] on select endpoints — SuperAdmin bypasses.",
        "Rate limiting: Redis-backed RateLimitingMiddleware when configured.",
        "SecurityHeadersMiddleware (HSTS, X-Frame-Options, CSP report-only).",
        "Swagger limited to Development/Staging only.",
        "File download IDOR fixed: IsVisibleToClient enforced in FileService.",
        "Remaining risks: JWT key in config (use env/Key Vault); InputSanitizationMiddleware disabled; Hangfire /hangfire in prod.",
    ])

    # 9
    h(doc, "9. Frontend System Design", 1)
    bullets(doc, [
        "Lazy feature modules; MainLayoutComponent single shell (no sidebar blink).",
        "Interceptors: CookieCredentials, Csrf, Token, HttpLoadingError (toasts only — no full-screen loader).",
        "Scoped loading UX: app-skeleton-dashboard on dashboard, orders, invoices, analytics.",
        "Quotes module at /quotes (Client, Admin, SuperAdmin).",
        "No NgRx — BehaviorSubject in services; shareReplay caches cleared on logout.",
        "PrimeNG only; paths @core, @shared, @environments.",
        "Gaps: PermissionGuard unused; large order-detail/dashboard components.",
    ])
    p(doc, "Key files: Frontend/src/app/app-routing.module.ts, core/services/auth.service.ts, layout/main-layout/main-layout.component.ts")

    # 10
    h(doc, "10. Backend System Design", 1)
    bullets(doc, [
        "Controllers thin → Application services → IApplicationDbContext.",
        "ScalabilityServiceRegistration: Redis cache, distributed rate limit, Hangfire on Redis, SignalR backplane.",
        "ReadModelCacheVersions + DistributedJsonCache for list/dashboard cache epochs.",
        "Hangfire via IBackgroundJobScheduler: EmailHangfireJobs, MaintenanceHangfireJobs, BillingAutoInvoiceService.",
        "SignalR: notification + entity update senders.",
        "DatabaseInitializationHostedService for migrations/seed at startup.",
        "ExceptionMiddleware maps exceptions to 400/401/403/404/500 JSON.",
        "NOT CQRS/MediatR; IRepository registered but unused.",
    ])

    # 11
    h(doc, "11. Environment & Configuration", 1)
    table(doc, ["Key", "Purpose"], [
        ["ConnectionStrings:DefaultConnection", "MySQL"],
        ["ConnectionStrings:Redis", "Cache, rate limit, SignalR, Hangfire (required when AllowInMemoryFallback=false)"],
        ["Scalability:AllowInMemoryFallback", "false in prod — mandates Redis"],
        ["AuthCookies:*", "HttpOnly cookies + CSRF cookie/header names"],
        ["Jwt:Key, Issuer, Audience", "JWT signing (min 32 chars)"],
        ["Cors:AllowedOrigins", "SPA origins (with credentials for cookies)"],
        ["FileStorage:Path", "Upload root (not web-served)"],
        ["Email:*", "SMTP + FrontendUrl for reset links"],
        ["ProductionSafety:*", "Kill-switches"],
        ["RateLimiting:*", "General, auth, orders, file upload per-minute"],
    ])
    p(doc, "Frontend: Frontend/src/environments/environment.prod.ts → apiUrl https://api.hawkmerchandising.com")
    p(doc, "IIS: Backend/src/LogoDesignPortal.API/web.config — 500MB upload, ASPNETCORE_ENVIRONMENT=Production")

    # 12
    h(doc, "12. Development Standards", 1)
    bullets(doc, [
        "C#: nullable, async services, I*Service interfaces, *Dto naming.",
        "Angular: feature folders, *-routing.module.ts, SCSS per component.",
        "Business logic in Application services; privacy masking in services when building DTOs.",
        "Order transitions via OrderStatusStateMachine only.",
    ])

    # 13
    h(doc, "13. Dependency Audit", 1)
    bullets(doc, [
        "Critical: EF Core + Pomelo MySQL, JWT, Hangfire, Redis, Serilog.",
        "Risk: Angular 15 EOL; legacy Backend duplicate folder; OpenTelemetry beta EF package.",
        "Coupling: Frontend DTO shapes tied to API camelCase JSON.",
    ])

    # 14
    h(doc, "14. Performance Analysis", 1)
    bullets(doc, [
        "GetAllOrders capped; paged endpoints for large datasets; analytics indexes in migrations.",
        "Redis in-memory fallback breaks multi-instance consistency if misconfigured.",
        "Large dashboard components; bundle budget 500kb warn / 1mb error.",
        "Local file storage bottleneck for horizontal scale — needs shared/blob storage.",
    ])

    # 15
    h(doc, "15. Technical Debt", 1)
    bullets(doc, [
        "Legacy Backend/LogoDesignPortal.API/ duplicate tree.",
        "Unused PermissionGuard; Permission system partially wired on controllers.",
        "Legacy price fields on LogoOrder (Price vs ClientChargePrice).",
        "InputSanitizationMiddleware disabled; Hangfire dashboard exposure in prod.",
        "Large order-detail and dashboard components.",
    ])

    # 16
    h(doc, "16. Testing Strategy", 1)
    table(doc, ["Layer", "Tool", "Location"], [
        ["Domain unit", "xUnit", "LogoDesignPortal.Domain.Tests"],
        ["Application unit", "xUnit + EF InMemory", "LogoDesignPortal.Application.Tests"],
        ["API integration", "WebApplicationFactory", "LogoDesignPortal.API.IntegrationTests"],
        ["Frontend unit", "Karma/Jasmine", "Frontend/src/**/*.spec.ts"],
        ["E2E", "Playwright", "Frontend/e2e/tests/"],
    ])
    p(doc, "CI: .github/workflows/test.yml — 80% backend line coverage; production build verification.")

    # 17
    h(doc, "17. Deployment & DevOps", 1)
    bullets(doc, [
        "No Docker in repo; IIS deploy API + static Angular dist.",
        "DB: EF migrations on startup or Backend/scripts/run-migration.bat.",
        "Environments: Development, Staging, Production, Testing (SQLite).",
        "Ops docs: docs/PRODUCTION-TROUBLESHOOTING.md, PRODUCTION-MIGRATION-SAFETY.md",
    ])

    # 18
    h(doc, "18. Developer Onboarding Guide", 1)
    p(doc, "Backend: cd Backend && dotnet restore && cd src/LogoDesignPortal.API && dotnet run")
    p(doc, "Frontend: cd Frontend && npm install && npm start")
    p(doc, "Default seed (rotate in prod): superadmin@logodesign.com / SuperAdmin@123")
    h(doc, "Add feature checklist", 2)
    bullets(doc, [
        "Domain → EF config → migration",
        "DTO + MappingProfile → I*Service → Controller with [Authorize]",
        "Angular service → component → route",
        "Apply role masking for cross-role data",
        "Tests: domain + application + integration",
    ])

    # 19
    h(doc, "19. AI Development Context", 1)
    bullets(doc, [
        "Use Backend/src/ only; never legacy folders.",
        "Backend must enforce security — not only Angular guards.",
        "Use OrderStatusStateMachine / OrderStatusTransitionHelper for status changes.",
        "Apply DTO masking for new client/designer fields.",
        "Register services in Application/DependencyInjection.cs.",
        "Respect ProductionSafetyOptions for billing/files/payout.",
        "Cookie auth: test mutations with CSRF header or Bearer in integration tests.",
        "Forbidden: exposing PasswordHash; static Files/; Material UI on frontend.",
    ])

    # 20
    h(doc, "20. Final Architecture Summary", 1)
    table(doc, ["Metric", "Assessment"], [
        ["Scalability", "Improved — Redis-backed rate limit, Hangfire, SignalR; shared file storage still needed"],
        ["Maintainability", "7/10 — consistent services; large components remain"],
        ["Production readiness", "7/10 — cookie auth + CSRF done; secrets in env, Redis mandatory"],
    ])
    h(doc, "Top recommendations", 2)
    bullets(doc, [
        "Rotate JWT/secrets via env/Key Vault only.",
        "Configure Redis in prod; set Scalability:AllowInMemoryFallback=false.",
        "Upgrade Angular LTS; decompose order-detail/dashboard.",
        "Wire PermissionGuard or remove; shared blob storage for files.",
        "Delete/archive legacy Backend duplicate tree; restrict Hangfire dashboard.",
    ])

    # Appendices
    h(doc, "Appendix A — Feature-to-file mapping", 1)
    table(doc, ["Feature", "Backend", "Frontend"], [
        ["Auth", "AuthController, AuthCookieService, AuthService", "auth/, interceptors, auth.service.ts"],
        ["Orders", "OrdersController, OrderService", "orders/*"],
        ["Quotes", "QuotesController, QuoteService", "quotes/*"],
        ["Files", "FilesController, FileService", "files/, order upload"],
        ["Billing", "BillingController, BillingService", "billing if routed"],
        ["Invoices", "InvoicesController, InvoiceService", "invoices/*"],
        ["Financial", "Admin/FinancialController", "financial/*"],
        ["Notifications", "NotificationHub", "realtime-notification.service.ts"],
    ])

    h(doc, "Appendix B — Critical path files", 1)
    bullets(doc, [
        "Backend/src/LogoDesignPortal.API/Program.cs",
        "Backend/src/LogoDesignPortal.API/Hosting/ScalabilityServiceRegistration.cs",
        "Backend/src/LogoDesignPortal.API/Services/AuthCookieService.cs",
        "Backend/src/LogoDesignPortal.Application/Services/OrderService.cs",
        "Backend/src/LogoDesignPortal.Domain/OrderStatusStateMachine.cs",
        "Backend/src/LogoDesignPortal.Infrastructure/Persistence/ApplicationDbContext.cs",
        "Frontend/src/app/app-routing.module.ts",
        "Frontend/src/app/core/services/auth.service.ts",
        "Frontend/src/app/app.module.ts (interceptor order)",
        ".github/workflows/test.yml",
        "docs/UPDATED_ARCHITECTURE_ONBOARDING.md",
        "docs/PRODUCTION-SAFETY-KILL-SWITCH.md",
    ])

    h(doc, "Appendix C — Read these files first (priority)", 1)
    p(doc, "Day 1: OrderStatus.cs, OrderStatusStateMachine.cs, LogoOrder.cs, app-routing.module.ts, main-layout.component.ts")
    p(doc, "Day 2: Program.cs, OrdersController.cs, OrderService.cs, auth.service.ts, token.interceptor.ts, order-detail.component.ts")
    p(doc, "Day 3: BillingService, InvoiceService, ApplicationDbContext.cs, TESTING_STANDARDS.md")
    p(doc, "Before production: appsettings.Production.json, web.config, PRODUCTION-TROUBLESHOOTING.md")

    try:
        doc.save(OUT)
        print(f"Created: {OUT}")
    except PermissionError:
        doc.save(OUT_FALLBACK)
        print(f"Could not overwrite {OUT} (file may be open). Created: {OUT_FALLBACK}")


if __name__ == "__main__":
    main()
